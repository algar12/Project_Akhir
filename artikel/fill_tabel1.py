#!/usr/bin/env python3
"""
fill_tabel1.py — Unified Network Guard
=======================================
Mengisi Tabel 1 di Artikel_IKOMTI_Unified_Network_Guard_bak2.docx dari
hasil ml/evaluator_tabel1.py (ml/models/tabel1_eval.json).

Tabel 1: 26 baris x 8 kolom
  Baris 0  = header
  Baris 1-4   = S1 × (Suricata, Rule, IF, Hibrida)
  Baris 5-8   = S2
  Baris 9-12  = S3
  Baris 13-16 = S4
  Baris 17-20 = S5
  Baris 21-24 = S6
  Baris 25    = Rata-rata × Hibrida (Usulan)
Kolom: 0=Skenario, 1=Metode, 2=Accuracy, 3=Precision, 4=Recall, 5=F1, 6=FPR, 7=Latensi(ms)

Nilai ditampilkan dalam persen (kecuali latensi ms). '—' = N/A.
"""
import json
import shutil
from pathlib import Path
from docx import Document

ARTIKEL = Path(__file__).resolve().parent / "Artikel_IKOMTI_Unified_Network_Guard_bak2.docx"
EVAL = Path(__file__).resolve().parent.parent / "ml" / "models" / "tabel1_eval.json"
BACKUP = ARTIKEL.with_suffix(".preeval.docx")

METHOD_KEYS = ["suricata", "rule", "ml", "hibrida"]
SCENARIO_IDS = ["S1", "S2", "S3", "S4", "S5", "S6"]


def fmt_pct(v):
    if v is None:
        return "—"
    return f"{v*100:.2f}"


def fmt_lat(v):
    if v is None:
        return "—"
    return f"{v:.1f}"


def cell_metrics(mm):
    """mm = dict metrik satu metode. Kembalikan list 6 nilai string untuk kolom 2-7."""
    return [
        fmt_pct(mm["accuracy"]),
        fmt_pct(mm["precision"]),
        fmt_pct(mm["recall"]),
        fmt_pct(mm["f1"]),
        fmt_pct(mm["fpr"]),
        fmt_lat(mm["latency_ms"]),
    ]


def main():
    if not EVAL.exists():
        print(f"[!] {EVAL} tidak ada. Jalankan ml/evaluator_tabel1.py dulu.")
        return
    data = json.loads(EVAL.read_text(encoding="utf-8"))

    # Peta hasil per scenario id
    by_sid = {r["scenario"]: r for r in data["scenarios"]}
    avg = data["average"]

    # Backup sekali
    if not BACKUP.exists():
        shutil.copy2(ARTIKEL, BACKUP)
        print(f"[i] Backup: {BACKUP.name}")

    doc = Document(str(ARTIKEL))
    if not doc.tables:
        print("[!] Tidak ada tabel di artikel.")
        return
    table = doc.tables[0]
    nrows = len(table.rows)
    if nrows < 26:
        print(f"[!] Tabel hanya {nrows} baris (perlu 26).")
        return

    # Isi baris per skenario (1-24)
    filled = 0
    for si, sid in enumerate(SCENARIO_IDS):
        if sid not in by_sid:
            print(f"[!] {sid} tidak ada di hasil evaluasi — baris dibiarkan '—'")
            continue
        res = by_sid[sid]
        for mi, m in enumerate(METHOD_KEYS):
            row_idx = 1 + si * 4 + mi
            mm = res["methods"][m]
            vals = cell_metrics(mm)
            cells = table.rows[row_idx].cells
            # kolom 0 & 1 sudah berisi Skenario & Metode (dari template)
            for ci, val in enumerate(vals):
                cells[2 + ci].text = val
            filled += 1

    # Baris 25 = Rata-rata × Hibrida (Usulan)
    avg_h = avg["hibrida"]
    vals = cell_metrics(avg_h)
    cells = table.rows[25].cells
    for ci, val in enumerate(vals):
        cells[2 + ci].text = val
    filled += 1

    doc.save(str(ARTIKEL))
    print(f"[OK] {filled} baris terisi di Tabel 1 → {ARTIKEL.name}")
    print("     Verifikasi: buka dokumen, cek Tabel 1.")
    print("     Jika salah, restore dari backup:", BACKUP.name)


if __name__ == "__main__":
    main()
